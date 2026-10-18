"""The results of a batch, next to the rows they belong to, in a copy of the table.

People who work from a sheet want it back with the outcome of every row. The copy gets two
columns, "Maptelo outcome" and "Maptelo message", after the last used column (or in
them, when the table already has them from an earlier run), and one more for each value the
runs read on the screens ("Maptelo read: saldo"): a process that looks something up hands it
back in its row. The original file is never changed.
"""

from __future__ import annotations

import csv
from collections.abc import Collection, Mapping, Sequence
from io import StringIO
from pathlib import Path
from typing import Any

from maptelo.data.secrets import COVERED, cover_cells, secret_columns
from maptelo.data.tables import WORKBOOKS, Table

COLUMNS = ("Maptelo outcome", "Maptelo message")
READ = "Maptelo read: "
"""The beginning of the header of a column with what the runs read („Maptelo read: saldo”)."""


def results_path(table: Table, out_dir: Path) -> Path:
    assert table.path is not None
    return out_dir / f"{table.path.stem}.results{table.path.suffix}"


def write_results(
    table: Table,
    latest: Mapping[int, Mapping[str, Any]],
    out_dir: Path,
    secret: Collection[str] = (),
    rows: Sequence[Sequence[int]] | None = None,
) -> Path:
    """``latest``: record number (1 = the first record) -> its latest result. ``rows``:
    which rows of the table each record is (an order and its lines), else one row each; each
    of them gets its record's result. The copy covers columns that hold secrets: by their
    headers, and ``secret`` (where secret inputs come from)."""
    of = rows if rows is not None else [[index] for index in range(len(table.numbers))]
    read = list(dict.fromkeys(n for entry in latest.values() for n in entry.get("read") or {}))
    titles = (*COLUMNS, *(READ + name for name in read))
    by_row = {
        table.numbers[index]: (
            str(entry["outcome"]),
            str(entry.get("message", "")),
            *(str((entry.get("read") or {}).get(name, "")) for name in read),
        )
        for number, entry in latest.items()
        if 0 < number <= len(of)
        for index in of[number - 1]
        if index < len(table.numbers)
    }
    out = results_path(table, out_dir)
    out.parent.mkdir(parents=True, exist_ok=True)
    assert table.path is not None
    if table.path.suffix.lower() in WORKBOOKS:
        _workbook(table, by_row, out, secret, titles)
    else:
        _csv(table, by_row, out, secret, titles)
    return out


def _workbook(
    table: Table,
    by_row: Mapping[int, tuple[str, ...]],
    out: Path,
    secret: Collection[str],
    titles: Sequence[str],
) -> None:
    from openpyxl import load_workbook

    assert table.path is not None
    workbook = load_workbook(table.path)  # formulas and formatting stay as they are
    sheet = workbook[table.sheet] if table.sheet else workbook.worksheets[0]
    for worksheet in workbook.worksheets:  # a password in another sheet travels just as well
        number = table.header_number if worksheet is sheet else 1
        cover_cells(worksheet, number, secret)
    header = [cell.value for cell in sheet[table.header_number]]
    columns = []
    for title in titles:
        if title in header:
            columns.append(header.index(title) + 1)
        else:
            columns.append(max(sheet.max_column, len(header), *columns) + 1)
            sheet.cell(row=table.header_number, column=columns[-1], value=title)
    for row, values in by_row.items():
        for column, value in zip(columns, values, strict=True):
            sheet.cell(row=row, column=column, value=value)
    workbook.save(out)


def _csv(
    table: Table,
    by_row: Mapping[int, tuple[str, ...]],
    out: Path,
    secret: Collection[str],
    titles: Sequence[str],
) -> None:
    assert table.path is not None
    text = table.path.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    try:
        dialect: Any = csv.Sniffer().sniff(lines[0], delimiters=",;\t")
    except (csv.Error, IndexError):
        dialect = csv.excel
    rows = list(csv.reader(StringIO(text), dialect=dialect))
    headers = [cell.strip() for cell in rows[table.header_number - 1]] if rows else []
    ours = [h for h in headers if h in COLUMNS or h.startswith(READ)]  # from an earlier run
    kept = [i for i, header in enumerate(headers) if header not in ours]
    hidden = {i for i, h in enumerate(headers) if h in secret_columns(headers, secret)}
    with out.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file, delimiter=dialect.delimiter)
        for number, row in enumerate(rows, start=1):
            if number == table.header_number:
                writer.writerow([*[row[i] for i in kept], *titles])
                continue
            if number > table.header_number:
                row = [COVERED if i in hidden and c.strip() else c for i, c in enumerate(row)]
            writer.writerow(
                [
                    *[row[i] if i < len(row) else "" for i in kept],
                    *by_row.get(number, ("",) * len(titles)),
                ]
            )
