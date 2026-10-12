"""Secrets in the sheets people work from: account passwords, PINs, tokens.

A copy of the sheet goes with each recording (so a model learns which column each value
comes from) and with the results of a batch (the sheet back, with each row's outcome). Copies
travel: into support bundles, repro packs and shared folders. So a column that holds secrets
is covered in every copy (benchmark: secret_in_sheet_copy), while the person's own file
is never changed. A column holds secrets when its header says so, or when the model takes a
secret input from it.
"""

from __future__ import annotations

import csv
import re
import shutil
from collections.abc import Collection, Iterable
from pathlib import Path
from typing import Any

from maptelo.data.tables import WORKBOOKS, cell_text, header_row

SECRET_HEADER = re.compile(
    r"has[łl]o|password|passw|\bpwd\b|\bpin\b|token|secret|sekret|api[ _-]?key|"
    r"klucz (?:api|prywatny)|private key|\bcvv\b|\bcvc\b|\botp\b",
    re.IGNORECASE,
)
COVERED = "••••••"


def secret_columns(headers: Iterable[str], also: Collection[str] = ()) -> list[str]:
    """The headers of columns that hold secrets: by their words, or named in ``also``."""
    return [h for h in headers if h and (h in also or SECRET_HEADER.search(h))]


def covered_copy(source: Path, target: Path, also: Collection[str] = ()) -> list[str]:
    """Copy a sheet with the cells of its secret columns covered, in every sheet of a
    workbook; returns the headers covered (nothing covered: a plain copy)."""
    target.parent.mkdir(parents=True, exist_ok=True)
    if source.suffix.lower() in WORKBOOKS:
        return _workbook(source, target, also)
    return _csv(source, target, also)


def cover_cells(worksheet: Any, header_number: int, also: Collection[str] = ()) -> list[str]:
    """Cover the secret columns of an open worksheet (openpyxl) below its header row."""
    header = [cell_text(cell.value) for cell in worksheet[header_number]]
    covered = secret_columns(header, also)
    for index, title in enumerate(header):
        if title not in covered:
            continue
        for (cell,) in worksheet.iter_rows(
            min_row=header_number + 1, min_col=index + 1, max_col=index + 1
        ):
            if cell.value not in (None, ""):
                cell.value = COVERED
    return covered


def _workbook(source: Path, target: Path, also: Collection[str]) -> list[str]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(source)  # formulas and formatting stay as they are
    except Exception:  # not a workbook after all: a copy cannot say more than the file
        shutil.copyfile(source, target)
        return []
    covered: list[str] = []
    for worksheet in workbook.worksheets:
        values = [[cell_text(c) for c in row] for row in worksheet.iter_rows(values_only=True)]
        start = header_row(values)
        if start is not None:
            covered += cover_cells(worksheet, start + 1, also)
    workbook.save(target)
    return covered


def _csv(source: Path, target: Path, also: Collection[str]) -> list[str]:
    text = source.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    if not lines:
        shutil.copyfile(source, target)
        return []
    try:
        dialect: Any = csv.Sniffer().sniff(lines[0], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(lines, dialect=dialect))
    headers = [cell.strip() for cell in rows[0]]
    covered = secret_columns(headers, also)
    hidden = {i for i, header in enumerate(headers) if header in covered}
    with target.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file, delimiter=dialect.delimiter)
        writer.writerow(rows[0])
        for row in rows[1:]:
            writer.writerow(
                [COVERED if i in hidden and cell.strip() else cell for i, cell in enumerate(row)]
            )
    return covered
