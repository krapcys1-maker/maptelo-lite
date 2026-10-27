"""Tables people work from: a sheet of a workbook or a CSV file, read as rows of text.

Cells keep what people see: whole numbers without ``.0`` (in a CSV file too, as a program
writes them: 604217359.0), dates as YYYY-MM-DD (with the time only when there is one),
formulas as their last computed value. In a sheet, the
headers are the row with the columns the process needs (or, when that is not known, the
first row with at least two values: a title above the table has one). Empty rows are not
records.
"""

from __future__ import annotations

import csv
import re
from collections.abc import Collection
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from io import StringIO
from pathlib import Path
from typing import Any

WORKBOOKS = (".xlsx", ".xlsm")


class TableError(Exception):
    """The file cannot be read as a table (no such sheet, not a workbook...)."""


@dataclass(frozen=True)
class Table:
    headers: list[str]
    rows: list[dict[str, str]]
    source: str
    """For messages: the file, and the sheet of a workbook."""
    path: Path | None = None
    sheet: str | None = None
    """The sheet the rows come from (a workbook's), by its name."""
    header_number: int = 1
    numbers: list[int] = field(default_factory=list)
    """Where each row is in the file: its row in the sheet, or its line in a CSV file."""


def cell_text(value: object, number_format: str = "") -> str:
    """A cell as people see it, as text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if (
        isinstance(value, int | float)
        and re.fullmatch(r"0+", number_format)
        and float(value).is_integer()
    ):
        return str(int(value)).zfill(len(number_format))
    if isinstance(value, datetime):
        if value.time() == time():
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(Decimal(repr(value)).normalize(), "f")
    return str(value).strip()


def _rows(
    headers: list[str], values: list[list[str]], first: int
) -> tuple[list[dict[str, str]], list[int]]:
    """The records (empty rows are not records) and where each one is (``first``: where
    ``values[0]`` is)."""
    named = [h or f"#{i + 1}" for i, h in enumerate(headers)]
    rows, numbers = [], []
    for offset, row in enumerate(values):
        if any(row):
            rows.append({name: (row[i] if i < len(row) else "") for i, name in enumerate(named)})
            numbers.append(first + offset)
    return rows, numbers


_FLOAT_WHOLE = re.compile(r"-?\d+\.0")
"""A whole number as a program writes a sheet's number into a CSV file (604217359.0): people
see 604217359, and a field takes it so."""


def _csv_cell(text: str) -> str:
    cell = text.strip()
    return cell[:-2] if _FLOAT_WHOLE.fullmatch(cell) else cell


def _read_csv(path: Path) -> Table:
    text = path.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    try:
        dialect: Any = csv.Sniffer().sniff(lines[0], delimiters=",;\t")
    except (csv.Error, IndexError):
        dialect = csv.excel
    parsed = [
        [_csv_cell(cell) for cell in row] for row in csv.reader(StringIO(text), dialect=dialect)
    ]
    if not parsed:
        return Table([], [], path.name, path)
    rows, numbers = _rows(parsed[0], parsed[1:], 2)
    return Table(parsed[0], rows, path.name, path, numbers=numbers)


def _header(values: list[list[str]], columns: Collection[str]) -> int | None:
    """Which row holds the headers (an index into ``values``)."""
    filled = [i for i, row in enumerate(values) if any(row)]
    if not filled:
        return None
    wanted = set(columns)
    if wanted:
        best = max(filled[:50], key=lambda i: (len(wanted & set(values[i])), -i))
        if wanted & set(values[best]):
            return best
    wide = [i for i in filled if sum(bool(cell) for cell in values[i]) >= 2]
    return wide[0] if wide else filled[0]


def header_row(values: list[list[str]], columns: Collection[str] = ()) -> int | None:
    """Which row of a sheet's values holds the headers (an index into ``values``)."""
    return _header(values, columns)


def _read_workbook(path: Path, sheet: str | None, columns: Collection[str]) -> Table:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several kinds for a broken or foreign file
        raise TableError(f"{path.name} cannot be read as a workbook: {exc}") from exc
    try:
        if sheet is None:
            worksheet = workbook.worksheets[0]
        elif sheet in workbook.sheetnames:
            worksheet = workbook[sheet]
        else:
            names = ", ".join(workbook.sheetnames)
            raise TableError(f"{path.name} has no sheet '{sheet}' (sheets: {names})")
        values = [
            [cell_text(c.value, c.number_format or "") for c in row]
            for row in worksheet.iter_rows(min_row=1)
        ]
        title = worksheet.title
    finally:
        workbook.close()
    source = f"{path.name}, sheet {title}"
    start = _header(values, columns)
    if start is None:
        return Table([], [], source, path, title)
    headers = values[start]
    while headers and not headers[-1]:  # trailing empty header cells
        headers = headers[:-1]
    data = [row[: len(headers)] for row in values[start + 1 :]]
    rows, numbers = _rows(headers, data, start + 2)
    return Table(headers, rows, source, path, title, start + 1, numbers)


def sheet_names(path: Path) -> list[str]:
    """The sheets of a workbook, in their order (a CSV file has none)."""
    from openpyxl import load_workbook

    if path.suffix.lower() not in WORKBOOKS:
        return []
    try:
        book = load_workbook(path, read_only=True)
    except Exception as exc:
        raise TableError(f"{path.name} cannot be read as a workbook: {exc}") from exc
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def read_table(path: Path, sheet: str | None = None, columns: Collection[str] = ()) -> Table:
    """A workbook (.xlsx) or a CSV file (comma, semicolon or tab) as a table; ``columns``
    are headers the process needs, which tell the header row apart from a title."""
    if path.suffix.lower() in WORKBOOKS:
        return _read_workbook(path, sheet, columns)
    if sheet is not None:
        raise TableError(f"{path.name} is not a workbook: it has no sheets")
    return _read_csv(path)
