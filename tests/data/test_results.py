"""The outcome of every row, in a copy of the table people gave."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from maptelo.data import COLUMNS, read_table, write_results

LATEST = {
    1: {"outcome": "VERIFIED", "message": "every effect confirmed by an independent read"},
    2: {"outcome": "BUSINESS_EXCEPTION", "message": "Klient nie ma produktu o podanym CPN"},
}


def sheet_file(path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Ceny"
    sheet.append(["Cennik od działu handlowego"])  # a title above the table
    sheet.append([])
    sheet.append(["Klient", "CPN", "Nowa cena"])
    sheet.append(["C001", "48212-B", 135])
    sheet.append([None, None, None])
    sheet.append(["C002", "99999-Z", 10, None, "uwaga poza tabelą"])
    sheet["C4"] = "=134+1"  # a formula stays a formula in the copy
    book.save(path)
    return path


def test_the_outcome_goes_next_to_its_row_in_a_copy(tmp_path: Path) -> None:
    source = sheet_file(tmp_path / "ceny.xlsx")
    table = read_table(source, "Ceny")
    assert (table.header_number, table.numbers) == (3, [4, 6])  # below the title
    assert table.headers == ["Klient", "CPN", "Nowa cena"]
    out = write_results(table, LATEST, tmp_path / "run")
    assert out == tmp_path / "run" / "ceny.results.xlsx"
    sheet = load_workbook(out)["Ceny"]
    header_row = [c.value for c in sheet[3]]
    assert header_row[-2:] == list(COLUMNS)
    outcome_column = header_row.index(COLUMNS[0]) + 1
    assert outcome_column == 6  # after everything used, so no note is overwritten
    assert sheet.cell(row=4, column=outcome_column).value == "VERIFIED"
    assert sheet.cell(row=5, column=outcome_column).value is None  # an empty row
    assert sheet.cell(row=6, column=outcome_column).value == "BUSINESS_EXCEPTION"
    assert sheet["C4"].value == "=134+1"
    assert load_workbook(source)["Ceny"].max_column == 5  # the original is not changed
    again = write_results(read_table(out, "Ceny"), LATEST, tmp_path / "run2")
    assert [c.value for c in load_workbook(again)["Ceny"][3]].count(COLUMNS[0]) == 1


def test_a_csv_gets_the_columns_too(tmp_path: Path) -> None:
    source = tmp_path / "ceny.csv"
    source.write_text("customer;cpn\nC001;48212-B\n\nC002;99999-Z\n", encoding="utf-8")
    out = write_results(read_table(source), LATEST, tmp_path)
    assert out.read_text(encoding="utf-8-sig").splitlines() == [
        "customer;cpn;Maptelo outcome;Maptelo message",
        "C001;48212-B;VERIFIED;every effect confirmed by an independent read",
        ";;;",
        "C002;99999-Z;BUSINESS_EXCEPTION;Klient nie ma produktu o podanym CPN",
    ]


def test_the_columns_a_process_needs_find_the_header_row(tmp_path: Path) -> None:
    """A two-cell title would pass for headers; the process's own columns do not lie."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Cennik", "październik 2026"])
    sheet.append(["Klient", "CPN", "Nowa cena"])
    sheet.append(["C001", "48212-B", 135])
    book.save(tmp_path / "ceny.xlsx")
    assert read_table(tmp_path / "ceny.xlsx").headers == ["Cennik", "październik 2026"]
    table = read_table(tmp_path / "ceny.xlsx", columns=["Klient", "Nowa cena"])
    assert (table.headers, table.numbers) == (["Klient", "CPN", "Nowa cena"], [3])
