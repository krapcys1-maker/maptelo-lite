"""Passwords in the sheet people work from stay out of every copy: the one that goes with a
recording and the one that comes back with the results (benchmark: secret_in_sheet_copy).
The person's own file is never changed."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from maptelo.data import read_table, write_results
from maptelo.data.secrets import COVERED, covered_copy, secret_columns

LATEST = {1: {"outcome": "VERIFIED", "message": "ok"}}


def accounts(path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Nowi klienci"
    sheet.append(["Konta dla nowych klientów"])  # a title above the table
    sheet.append(["Login", "Imię", "Hasło", "Klucz"])
    sheet.append(["anowak", "Anna", "Tajne!2026", "k-1"])
    sheet.append(["pkowal", "Piotr", "", "k-2"])
    other = book.create_sheet("Dostępy")
    other.append(["System", "PIN"])
    other.append(["bank", "4321"])
    book.save(path)
    return path


def test_secret_columns_are_known_by_their_words() -> None:
    headers = ["Login", "Hasło", "haslo tymczasowe", "Password", "PIN", "API key", "Token SMS"]
    assert secret_columns(headers) == headers[1:]
    assert secret_columns(["Klient", "Kod pocztowy", "Opinia", "Pinezka"]) == []
    assert secret_columns(["Klucz", "Klient"], also=["Klucz"]) == ["Klucz"]


def test_the_copy_of_a_workbook_covers_passwords_in_every_sheet(tmp_path: Path) -> None:
    source = accounts(tmp_path / "klienci.xlsx")
    copy = tmp_path / "rec" / "data" / "klienci.xlsx"
    assert covered_copy(source, copy) == ["Hasło", "PIN"]
    book = load_workbook(copy)
    rows = [[c.value for c in row] for row in book["Nowi klienci"].iter_rows(min_row=2)]
    assert rows == [
        ["Login", "Imię", "Hasło", "Klucz"],
        ["anowak", "Anna", COVERED, "k-1"],
        ["pkowal", "Piotr", None, "k-2"],  # an empty cell stays empty
    ]
    assert [c.value for c in book["Dostępy"][2]] == ["bank", COVERED]
    original = load_workbook(source)["Nowi klienci"]
    assert original["C3"].value == "Tajne!2026"  # the person's file is not changed
    table = read_table(copy, "Nowi klienci")  # the copy is still a table to learn from
    assert table.headers == ["Login", "Imię", "Hasło", "Klucz"]
    assert table.rows[0]["Login"] == "anowak"


def test_the_copy_of_a_csv_covers_passwords(tmp_path: Path) -> None:
    source = tmp_path / "klienci.csv"
    source.write_text("Login;Hasło;Imię\nanowak;Tajne!2026;Anna\n", encoding="utf-8")
    copy = tmp_path / "copy.csv"
    assert covered_copy(source, copy) == ["Hasło"]
    assert copy.read_text(encoding="utf-8-sig").splitlines() == [
        "Login;Hasło;Imię",
        f"anowak;{COVERED};Anna",
    ]


def test_the_results_cover_passwords_and_the_columns_of_secret_inputs(tmp_path: Path) -> None:
    source = accounts(tmp_path / "klienci.xlsx")
    table = read_table(source, "Nowi klienci")
    out = write_results(table, LATEST, tmp_path / "run", secret=["Klucz"])
    book = load_workbook(out)
    row = [c.value for c in book["Nowi klienci"][3]]
    assert row[:4] == ["anowak", "Anna", COVERED, COVERED]
    assert row[-2:] == ["VERIFIED", "ok"]
    assert [c.value for c in book["Dostępy"][2]] == ["bank", COVERED]

    csv_source = tmp_path / "klienci.csv"
    csv_source.write_text("Login;Hasło\nanowak;Tajne!2026\n", encoding="utf-8")
    csv_out = write_results(read_table(csv_source), LATEST, tmp_path / "run-csv")
    assert (
        csv_out.read_text(encoding="utf-8-sig").splitlines()[1] == f"anowak;{COVERED};VERIFIED;ok"
    )


def test_a_sheet_that_cannot_be_opened_is_not_copied_in_clear(tmp_path) -> None:
    from maptelo.data.secrets import CoverError

    source = tmp_path / "hasla.xlsx"
    source.write_bytes(b"not a workbook: Haslo;admin123")
    copy = tmp_path / "copy" / "hasla.xlsx"
    with pytest.raises(CoverError):
        covered_copy(source, copy)
    assert not copy.exists()  # no copy whose secret columns stayed in clear


def test_a_java_recording_keeps_its_sheet_with_passwords_covered(tmp_path: Path) -> None:
    from maptelo.capture.java import JavaRecorder, JavaRecorderOptions

    recorder = JavaRecorder.__new__(JavaRecorder)
    recorder.options = JavaRecorderOptions(data=accounts(tmp_path / "klienci.xlsx"))
    recorder.folder = tmp_path / "rec"
    recorder.folder.mkdir()
    kept = recorder._keep_data()
    assert kept is not None
    assert kept.file == "data.xlsx"
    rows = list(load_workbook(recorder.folder / "data.xlsx").worksheets[0].values)
    assert "Tajne!2026" not in str(rows)
    assert COVERED in str(rows)
